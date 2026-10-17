"""Reads the text of a Word document (.docx) with python-docx: its paragraphs, then its tables.

The older .doc format is a different, binary format that python-docx cannot open, so it is not accepted.

Typical usage example:

  extracted = WordTextReader().read(Path('plan.docx'))
"""

import zipfile
from pathlib import Path

import docx
import docx.opc.exceptions

from trading_agent.attachments import text_reader


class WordTextReader(text_reader.TextReader):
    """Reads a .docx file."""

    extensions = ('.docx',)
    media_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

    def read(self, path: Path) -> text_reader.ExtractedText:
        """Reads the paragraphs, and each table row with its cells separated by " | ".

        Args:
            path (Path): The document.

        Returns:
            text_reader.ExtractedText: Its text.

        Raises:
            TextReadError: The file is damaged or is not a .docx file.
        """
        try:
            document = docx.Document(str(path))
        except (
            docx.opc.exceptions.PackageNotFoundError,
            zipfile.BadZipFile,
            KeyError,
            ValueError,
        ) as error:
            raise text_reader.TextReadError(
                f'the Word document cannot be read ({type(error).__name__}: {error})'
            ) from error
        lines = []
        for paragraph in document.paragraphs:
            lines.append(paragraph.text)
        for number, table in enumerate(document.tables, start=1):
            lines.append('')
            lines.append(f'--- Table {number} ---')
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(cell.text.strip())
                lines.append(' | '.join(cells))
        return text_reader.ExtractedText(text='\n'.join(lines).strip())
