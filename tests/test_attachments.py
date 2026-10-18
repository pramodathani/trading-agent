"""Tests for pictures and documents sent with messages, from the upload to the model."""

import asyncio
import io
import json
from pathlib import Path

import docx
import fastapi
import fastapi.testclient
import openpyxl
import PIL.Image

from tests import test_claude_code_backend
from tests import test_openai_compatible_backend
from trading_agent.agent import agent_loop
from trading_agent.agent import conversation as conversation_module
from trading_agent.agent import messages
from trading_agent.attachments import (
    attachment_store as attachment_store_module,
)
from trading_agent.attachments import image_preparer as image_preparer_module
from trading_agent.backends import ollama_backend
from trading_agent.channels import notification_formatter
from trading_agent.channels import pairing_store as pairing_store_module
from trading_agent.channels import telegram_api
from trading_agent.channels import telegram_channel
from trading_agent.configuration import agent_configuration
from trading_agent.gateway import access_guard
from trading_agent.gateway.routes import attachment_routes
from trading_agent.gateway.routes import desk_routes
from trading_agent.runtime import desk_runtime
from trading_agent.runtime import event_bus as event_bus_module
from trading_agent.runtime import trigger as trigger_module
from trading_agent.runtime import turn_runner
from trading_agent.workspace import workspace_layout

_TOKEN = 'test-access-token'
_OWNER = 111
_STRANGER = 999


def _picture(
    width: int = 40,
    height: int = 30,
    mode: str = 'RGB',
    picture_format: str = 'PNG',
) -> bytes:
    """Draws a plain picture.

    Args:
        width (int): Its width in pixels.
        height (int): Its height in pixels.
        mode (str): Pillow's colour mode, such as RGB or RGBA.
        picture_format (str): The file format to save it in.

    Returns:
        bytes: The encoded picture.
    """
    output = io.BytesIO()
    PIL.Image.new(mode, (width, height), 'orange').save(
        output, format=picture_format
    )
    return output.getvalue()


def _text_pdf(text: str) -> bytes:
    """Writes a one-page PDF whose text layer holds one line.

    Args:
        text (str): The line.

    Returns:
        bytes: The PDF.
    """
    stream = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode('latin-1')
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',
        b'<< /Length '
        + str(len(stream)).encode()
        + b' >>\nstream\n'
        + stream
        + b'\nendstream',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    output = io.BytesIO()
    output.write(b'%PDF-1.4\n')
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f'{number} 0 obj\n'.encode() + body + b'\nendobj\n')
    cross_reference = output.tell()
    output.write(f'xref\n0 {len(objects) + 1}\n'.encode())
    output.write(b'0000000000 65535 f \n')
    for offset in offsets:
        output.write(f'{offset:010d} 00000 n \n'.encode())
    output.write(
        f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{cross_reference}\n%%EOF\n'.encode()
    )
    return output.getvalue()


def _word_document() -> bytes:
    """Writes a .docx with a paragraph and a table.

    Returns:
        bytes: The document.
    """
    document = docx.Document()
    document.add_paragraph('Buy RELIANCE below 2400.')
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'Symbol'
    table.rows[0].cells[1].text = 'Target'
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _workbook() -> bytes:
    """Writes a .xlsx with one sheet of two rows.

    Returns:
        bytes: The workbook.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Positions'
    sheet.append(
        [
            'Symbol',
            'Quantity',
        ]
    )
    sheet.append(
        [
            'INFY',
            10,
        ]
    )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _store(directory: Path) -> attachment_store_module.AttachmentStore:
    """Builds a store over a scratch workspace with a main desk.

    Args:
        directory (Path): A scratch folder.

    Returns:
        attachment_store_module.AttachmentStore: The store.
    """
    return attachment_store_module.AttachmentStore(
        workspace_layout.WorkspaceLayout(directory)
    )


class TestAttachmentStore:
    """Tests for keeping uploads and reading documents."""

    def test_picture_is_stored_and_found_again(self, tmp_path: Path) -> None:
        """A picture is kept in its own folder and found again by its identifier."""
        store = _store(tmp_path)

        attachment = store.save('main', 'chart.png', _picture())
        found = store.find('main', attachment.attachment_id)

        assert found == attachment
        assert attachment.is_image
        assert attachment.media_type == 'image/png'
        assert attachment.desk_path.startswith('uploads/')
        assert attachment.desk_path.endswith(
            f'/{attachment.attachment_id}/chart.png'
        )
        assert attachment.stored_file().read_bytes() == _picture()

    def test_picture_type_comes_from_its_contents(self, tmp_path: Path) -> None:
        """A JPEG named .png is stored as the JPEG it is."""
        attachment = _store(tmp_path).save(
            'main', 'photo.png', _picture(picture_format='JPEG')
        )

        assert attachment.media_type == 'image/jpeg'

    def test_fake_picture_is_refused_and_removed(self, tmp_path: Path) -> None:
        """Text named .png is refused, and nothing of it is left behind."""
        store = _store(tmp_path)

        try:
            store.save('main', 'chart.png', b'not a picture')
        except attachment_store_module.AttachmentRefusedError as error:
            assert 'chart.png' in str(error)
        else:
            raise AssertionError('The fake picture was accepted')
        uploads = tmp_path / 'desks' / 'main' / 'uploads'
        assert list(uploads.rglob('*.png')) == []

    def test_empty_large_and_unknown_files_are_refused(
        self, tmp_path: Path
    ) -> None:
        """Empty files, files over 20 MB and programs are refused with the reason."""
        store = _store(tmp_path)
        cases = [
            (
                'empty.txt',
                b'',
                'is empty',
            ),
            (
                'huge.txt',
                b'x' * (attachment_store_module.MAXIMUM_BYTES + 1),
                'the most a desk takes is 20 MB',
            ),
            (
                'tool.exe',
                b'MZ',
                'not a type the desks can use',
            ),
        ]
        for file_name, data, reason in cases:
            try:
                store.save('main', file_name, data)
            except attachment_store_module.AttachmentRefusedError as error:
                assert reason in str(error)
            else:
                raise AssertionError(f'{file_name} was accepted')

    def test_csv_text_is_labelled_as_outside_data(self, tmp_path: Path) -> None:
        """A document's text file is wrapped in the same markers as a web page."""
        attachment = _store(tmp_path).save(
            'main', 'trades.csv', b'symbol,quantity\nINFY,10\n'
        )

        text_file = tmp_path / 'desks' / 'main' / attachment.text_desk_path
        text = text_file.read_text(encoding='utf-8')
        assert text.startswith(
            '[Content from the uploaded file trades.csv begins.'
        )
        assert 'INFY,10' in text
        assert text.rstrip().endswith(
            '[Content from the uploaded file trades.csv ends.]'
        )
        assert attachment.text_characters == len('symbol,quantity\nINFY,10\n')

    def test_pdf_word_and_excel_are_read(self, tmp_path: Path) -> None:
        """PDF, Word and Excel files become text, with a PDF's page count."""
        store = _store(tmp_path)

        pdf = store.save('main', 'report.pdf', _text_pdf('Hello from the PDF'))
        word = store.save('main', 'plan.docx', _word_document())
        excel = store.save('main', 'book.xlsx', _workbook())

        desk = tmp_path / 'desks' / 'main'
        assert pdf.page_count == 1
        assert 'Hello from the PDF' in (desk / pdf.text_desk_path).read_text()
        word_text = (desk / word.text_desk_path).read_text()
        assert 'Buy RELIANCE below 2400.' in word_text
        assert 'Symbol | Target' in word_text
        excel_text = (desk / excel.text_desk_path).read_text()
        assert '--- Sheet Positions ---' in excel_text
        assert 'INFY\t10' in excel_text

    def test_unreadable_document_is_kept_with_the_reason(
        self, tmp_path: Path
    ) -> None:
        """A damaged Word file is kept, and the desk is told why it cannot read it."""
        attachment = _store(tmp_path).save('main', 'broken.docx', b'garbage')

        assert attachment.text_desk_path is None
        assert 'cannot be read' in attachment.problem
        assert 'Its text could not be read' in attachment.describe_for_model()

    def test_pdf_without_text_says_to_send_pictures(
        self, tmp_path: Path
    ) -> None:
        """A PDF holding only a picture, as a scan does, is reported as having no text."""
        output = io.BytesIO()
        PIL.Image.new('RGB', (100, 100), 'white').save(output, format='PDF')

        attachment = _store(tmp_path).save(
            'main', 'scan.pdf', output.getvalue()
        )

        assert attachment.text_desk_path is None
        assert 'send the pages as pictures' in attachment.problem

    def test_folders_in_the_name_are_dropped(self, tmp_path: Path) -> None:
        """A name with folders in it cannot place the file outside its upload folder."""
        attachment = _store(tmp_path).save(
            'main', '../../shared/TradingRules.md', b'# rules'
        )

        assert attachment.desk_path.endswith(
            f'/{attachment.attachment_id}/TradingRules.md'
        )
        assert not (tmp_path / 'shared' / 'TradingRules.md').exists()

    def test_unknown_identifier_is_not_found(self, tmp_path: Path) -> None:
        """An identifier that is not a stored upload, or not shaped like one, is not found."""
        store = _store(tmp_path)
        for attachment_id in ('0123456789ab', '../etc'):
            try:
                store.find('main', attachment_id)
            except attachment_store_module.AttachmentNotFoundError:
                continue
            raise AssertionError(f'{attachment_id} was found')


class TestImagePreparer:
    """Tests for preparing pictures for the models."""

    def test_large_picture_is_shrunk_to_a_jpeg(self, tmp_path: Path) -> None:
        """A 4000 by 3000 photo reaches the model at most 1568 pixels wide, as a JPEG."""
        path = tmp_path / 'photo.png'
        path.write_bytes(_picture(4000, 3000))

        image = image_preparer_module.ImagePreparer().prepare(path)

        assert image.media_type == 'image/jpeg'
        with PIL.Image.open(io.BytesIO(image.data)) as prepared:
            assert prepared.size == (1568, 1176)

    def test_transparent_picture_stays_a_png(self, tmp_path: Path) -> None:
        """A picture with transparency is sent as a PNG, which keeps it."""
        path = tmp_path / 'logo.png'
        path.write_bytes(_picture(mode='RGBA'))

        image = image_preparer_module.ImagePreparer().prepare(path)

        assert image.media_type == 'image/png'

    def test_missing_picture_is_left_out(self, tmp_path: Path) -> None:
        """A picture whose file is gone is skipped rather than failing the turn."""
        assert (
            image_preparer_module.ImagePreparer().prepare(tmp_path / 'gone.png')
            is None
        )


class _RecordingRunner(turn_runner.TurnRunner):
    """A turn runner that keeps every trigger it is given.

    Attributes:
        triggers: The triggers, in order.
    """

    def __init__(self):
        """Creates the runner."""
        self.triggers: list[trigger_module.Trigger] = []

    async def run_turn(
        self, trigger: trigger_module.Trigger
    ) -> turn_runner.TurnResult:
        """Keeps the trigger and answers briefly.

        Args:
            trigger (trigger_module.Trigger): The reason for the turn.

        Returns:
            turn_runner.TurnResult: A short answer.
        """
        self.triggers.append(trigger)
        return turn_runner.TurnResult(text='seen')


async def _run_one(
    desk: desk_runtime.DeskRuntime,
    runner: _RecordingRunner,
    trigger: trigger_module.Trigger,
) -> None:
    """Submits a trigger and runs the desk until the runner has it.

    Args:
        desk (desk_runtime.DeskRuntime): The desk.
        runner (_RecordingRunner): Its runner.
        trigger (trigger_module.Trigger): The trigger.
    """
    task = asyncio.create_task(desk.run())
    desk.submit(trigger)
    try:
        for _ in range(200):
            if runner.triggers:
                return
            await asyncio.sleep(0.01)
        raise AssertionError('The desk never ran the turn')
    finally:
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)


class TestMessagesWithAttachments:
    """Tests for how attachments travel with a trigger to the model."""

    def test_trigger_records_its_attachments(self, tmp_path: Path) -> None:
        """A trigger's dictionary lists its attachments, so the chat can show them after a reload."""
        attachment = _store(tmp_path).save('main', 'chart.png', _picture())
        trigger = trigger_module.Trigger.now(
            trigger_module.TriggerSource.USER_MESSAGE,
            'main',
            'What is this?',
            (attachment,),
        )

        dictionary = json.loads(json.dumps(trigger.to_json_dictionary()))

        assert dictionary['attachments'][0]['attachment_id'] == (
            attachment.attachment_id
        )
        assert trigger.images == (attachment,)

    def test_desk_adds_a_note_for_each_attachment(self, tmp_path: Path) -> None:
        """The turn's text names each attachment and says where a document's text is."""
        store = _store(tmp_path)
        picture = store.save('main', 'chart.png', _picture())
        document = store.save('main', 'trades.csv', b'symbol\nINFY\n')
        runner = _RecordingRunner()
        desk = desk_runtime.DeskRuntime(
            agent_configuration.DeskConfiguration(name='main'),
            event_bus_module.EventBus(),
            runner,
        )
        trigger = trigger_module.Trigger.now(
            trigger_module.TriggerSource.USER_MESSAGE,
            'main',
            'Look at these.',
            (
                picture,
                document,
            ),
        )

        asyncio.run(_run_one(desk, runner, trigger))

        text = runner.triggers[0].text
        assert 'Look at these.' in text
        assert (
            f'[Picture attached: chart.png (1 KB), saved at {picture.desk_path}.'
            in text
        )
        assert f'Its text is at {document.text_desk_path}' in text
        assert runner.triggers[0].attachments == trigger.attachments

    def test_message_with_attachments_never_steers(
        self, tmp_path: Path
    ) -> None:
        """A message with a picture waits for its own turn instead of joining the running one."""
        attachment = _store(tmp_path).save('main', 'chart.png', _picture())
        desk = desk_runtime.DeskRuntime(
            agent_configuration.DeskConfiguration(name='main'),
            event_bus_module.EventBus(),
            _RecordingRunner(),
        )
        desk._running_trigger_id = 'running'

        desk.submit(
            trigger_module.Trigger.now(
                trigger_module.TriggerSource.USER_MESSAGE,
                'main',
                'And this one?',
                (attachment,),
            )
        )

        assert desk.queued_count == 1
        assert desk.steering_inbox.take_text() is None

    def test_agent_loop_puts_pictures_on_the_user_message(
        self, tmp_path: Path
    ) -> None:
        """The agent loop sends the prepared pictures with the turn's user message."""
        attachment = _store(tmp_path).save('main', 'chart.png', _picture())
        loop = agent_loop.AgentLoop(
            'main',
            backend=None,
            executor=None,
            event_bus=event_bus_module.EventBus(),
            maximum_rounds=3,
        )
        trigger = trigger_module.Trigger.now(
            trigger_module.TriggerSource.USER_MESSAGE,
            'main',
            'What is this?',
            (attachment,),
        )

        images = loop._prepared_images(trigger)

        assert len(images) == 1
        assert images[0].media_type == 'image/jpeg'

    def test_ollama_and_openai_receive_the_pictures(self) -> None:
        """Ollama gets base64 pictures in "images"; OpenAI gets a text part and an image part."""
        image = messages.MessageImage('image/jpeg', b'\xff\xd8picture')
        conversation = conversation_module.Conversation('be brief')
        conversation.append(
            messages.ConversationMessage.user('What is this?', (image,))
        )
        ollama = ollama_backend.OllamaBackend(
            'gemma',
            agent_configuration.BackendConfiguration(
                kind='ollama',
                base_url='http://127.0.0.1:11434',
                model='gemma4:31b',
                context_window=32768,
            ),
        )

        ollama_wire = ollama._wire_messages(conversation)
        openai_wire = test_openai_compatible_backend._backend()._wire_messages(
            conversation
        )

        assert ollama_wire[1]['images'] == [
            image.base64_data(),
        ]
        assert openai_wire[1]['content'] == [
            {
                'type': 'text',
                'text': 'What is this?',
            },
            {
                'type': 'image_url',
                'image_url': {
                    'url': f'data:image/jpeg;base64,{image.base64_data()}',
                },
            },
        ]

    def test_claude_receives_a_message_of_text_and_pictures(
        self, tmp_path: Path
    ) -> None:
        """Claude is sent one user message whose content is the text and an image block."""
        attachment = _store(tmp_path).save('main', 'chart.png', _picture())
        runner = test_claude_code_backend._runner(event_bus_module.EventBus())
        trigger = trigger_module.Trigger.now(
            trigger_module.TriggerSource.USER_MESSAGE,
            'main',
            'What is this?',
            (attachment,),
        )
        prompt = test_claude_code_backend._FixedPromptBuilder().build()

        async def first_message() -> dict:
            """Reads the one message the query yields.

            Returns:
                dict: The message.
            """
            query = runner._query(prompt, trigger)
            sent = []
            async for message in query:
                sent.append(message)
            assert len(sent) == 1
            return sent[0]

        message = asyncio.run(first_message())

        content = message['message']['content']
        assert content[0]['type'] == 'text'
        assert content[0]['text'].endswith('What is this?')
        assert content[1]['type'] == 'image'
        assert content[1]['source']['media_type'] == 'image/jpeg'

    def test_claude_receives_plain_text_without_pictures(self) -> None:
        """Without pictures, Claude is sent the text alone, as before."""
        runner = test_claude_code_backend._runner(event_bus_module.EventBus())
        prompt = test_claude_code_backend._FixedPromptBuilder().build()
        runner._session_prompt_fingerprint = prompt.fingerprint

        query = runner._query(prompt, test_claude_code_backend._trigger())

        assert query == 'What time is it?'

    def test_pictures_count_toward_the_context_estimate(self) -> None:
        """Each picture adds 1500 tokens to the estimate compression watches."""
        image = messages.MessageImage('image/jpeg', b'x')
        conversation = conversation_module.Conversation('be brief')
        conversation.append(
            messages.ConversationMessage.user(
                '',
                (
                    image,
                    image,
                ),
            )
        )

        assert conversation.estimated_tokens() == 3000


class _StandInDesk:
    """A desk that keeps what it is sent.

    Attributes:
        triggers: The triggers submitted.
    """

    def __init__(self):
        """Creates the desk."""
        self.triggers: list[trigger_module.Trigger] = []

    def submit(self, trigger: trigger_module.Trigger) -> None:
        """Keeps a trigger.

        Args:
            trigger (trigger_module.Trigger): The trigger.
        """
        self.triggers.append(trigger)


def _client(
    directory: Path,
    desk: _StandInDesk,
) -> fastapi.testclient.TestClient:
    """Mounts the desk and attachment routes for one stand-in desk.

    Args:
        directory (Path): The scratch workspace.
        desk (_StandInDesk): The desk.

    Returns:
        fastapi.testclient.TestClient: A client sending the token.
    """
    layout = workspace_layout.WorkspaceLayout(directory)
    store = attachment_store_module.AttachmentStore(layout)
    guard = access_guard.AccessGuard(_TOKEN)
    desks = {
        'main': desk,
    }
    application = fastapi.FastAPI()
    application.include_router(
        desk_routes.DeskRoutes(desks, guard, attachment_store=store).router()
    )
    application.include_router(
        attachment_routes.AttachmentRoutes(desks, store, layout, guard).router()
    )
    return fastapi.testclient.TestClient(
        application,
        headers={
            'Authorization': f'Bearer {_TOKEN}',
        },
    )


class TestAttachmentRoutes:
    """Tests for uploading files and sending them with a message."""

    def test_upload_then_send_with_the_message(self, tmp_path: Path) -> None:
        """An uploaded picture is named in the message and reaches the desk with it."""
        desk = _StandInDesk()
        client = _client(tmp_path, desk)

        uploaded = client.post(
            '/api/desks/main/attachments',
            files={
                'file': (
                    'chart.png',
                    _picture(),
                    'image/png',
                ),
            },
        )
        attachment_id = uploaded.json()['attachment_id']
        sent = client.post(
            '/api/desks/main/messages',
            json={
                'text': '',
                'attachment_ids': [
                    attachment_id,
                ],
            },
        )

        assert uploaded.status_code == 201
        assert sent.status_code == 202
        assert desk.triggers[0].attachments[0].attachment_id == attachment_id

    def test_uploaded_picture_can_be_fetched(self, tmp_path: Path) -> None:
        """The chat page can fetch an uploaded picture to show it."""
        store = _store(tmp_path)
        attachment = store.save('main', 'chart.png', _picture())
        client = _client(tmp_path, _StandInDesk())

        answer = client.get(
            f'/api/desks/main/attachments/{attachment.attachment_id}'
        )

        assert answer.status_code == 200
        assert answer.headers['content-type'] == 'image/png'
        assert answer.content == _picture()

    def test_refused_upload_explains_why(self, tmp_path: Path) -> None:
        """A program is refused with 422 and the reason."""
        client = _client(tmp_path, _StandInDesk())

        answer = client.post(
            '/api/desks/main/attachments',
            files={
                'file': (
                    'tool.exe',
                    b'MZ',
                    'application/octet-stream',
                ),
            },
        )

        assert answer.status_code == 422
        assert 'not a type the desks can use' in answer.json()['detail']

    def test_empty_message_and_unknown_upload_are_refused(
        self, tmp_path: Path
    ) -> None:
        """A message with neither text nor files, or naming an upload that does not exist, is refused."""
        desk = _StandInDesk()
        client = _client(tmp_path, desk)

        empty = client.post('/api/desks/main/messages', json={'text': '  '})
        unknown = client.post(
            '/api/desks/main/messages',
            json={
                'text': 'see this',
                'attachment_ids': [
                    '0123456789ab',
                ],
            },
        )

        assert empty.status_code == 422
        assert unknown.status_code == 422
        assert desk.triggers == []

    def test_files_need_the_token(self, tmp_path: Path) -> None:
        """Uploads cannot be fetched without the access token."""
        store = _store(tmp_path)
        attachment = store.save('main', 'chart.png', _picture())
        client = _client(tmp_path, _StandInDesk())

        answer = client.get(
            f'/api/desks/main/attachments/{attachment.attachment_id}',
            headers={
                'Authorization': 'Bearer wrong',
            },
        )

        assert answer.status_code == 401


class _StandInTelegram:
    """Telegram's servers: messages that arrive, files to download, and messages sent.

    Attributes:
        sent: Every message sent, as (chat id, text).
        downloads: The file ids downloaded.
    """

    def __init__(self, files: dict[str, bytes]):
        """Creates the stand-in.

        Args:
            files (dict[str, bytes]): The files Telegram holds, by file id.
        """
        self._files = files
        self._incoming: asyncio.Queue = asyncio.Queue()
        self.sent: list[tuple[int, str]] = []
        self.menu: list[tuple[str, str]] = []
        self.downloads: list[str] = []

    async def start(self) -> str:
        """Pretends the token works.

        Returns:
            str: The bot's username.
        """
        return 'test_bot'

    async def messages(self) -> list[telegram_api.IncomingMessage]:
        """Waits for the next message.

        Returns:
            list[telegram_api.IncomingMessage]: One message.
        """
        return [
            await self._incoming.get(),
        ]

    async def download(self, file_id: str) -> bytes:
        """Hands over a file.

        Args:
            file_id (str): The file's id.

        Returns:
            bytes: Its contents.
        """
        self.downloads.append(file_id)
        return self._files[file_id]

    async def send(self, chat_id: int, text: str) -> None:
        """Records a sent message.

        Args:
            chat_id (int): The chat.
            text (str): The message.
        """
        self.sent.append((chat_id, text))

    async def set_commands(self, commands: tuple[tuple[str, str], ...]) -> None:
        """Records the command menu.

        Args:
            commands (tuple[tuple[str, str], ...]): Each command as (name, description).
        """
        self.menu = list(commands)

    async def close(self) -> None:
        """Does nothing."""

    def receive(self, message: telegram_api.IncomingMessage) -> None:
        """Makes a message arrive.

        Args:
            message (telegram_api.IncomingMessage): The message.
        """
        self._incoming.put_nowait(message)


async def _run_channel(
    channel: telegram_channel.TelegramChannel,
    desk: desk_runtime.DeskRuntime,
    api: _StandInTelegram,
    chat_id: int,
) -> None:
    """Runs the channel and the desk until a chat has had one reply.

    Args:
        channel (telegram_channel.TelegramChannel): The channel.
        desk (desk_runtime.DeskRuntime): The desk.
        api (_StandInTelegram): The stand-in.
        chat_id (int): The chat to watch.
    """
    tasks = [
        asyncio.create_task(desk.run()),
        asyncio.create_task(channel.run()),
    ]
    try:
        for _ in range(300):
            if any(chat == chat_id for chat, _text in api.sent):
                return
            await asyncio.sleep(0.01)
        raise AssertionError(f'Chat {chat_id} got {api.sent}')
    finally:
        for task in tasks:
            task.cancel()
        await asyncio.gather(*tasks, return_exceptions=True)


class TestTelegramAttachments:
    """Tests for photos and files sent to the bot."""

    def _channel(
        self, directory: Path, api: _StandInTelegram
    ) -> tuple[
        telegram_channel.TelegramChannel,
        desk_runtime.DeskRuntime,
        _RecordingRunner,
    ]:
        """Builds a channel with the owner paired to the main desk.

        Args:
            directory (Path): A scratch folder.
            api (_StandInTelegram): The stand-in.

        Returns:
            tuple[telegram_channel.TelegramChannel, desk_runtime.DeskRuntime, _RecordingRunner]: The channel, its desk and the desk's runner.
        """
        event_bus = event_bus_module.EventBus()
        runner = _RecordingRunner()
        desk = desk_runtime.DeskRuntime(
            agent_configuration.DeskConfiguration(name='main'),
            event_bus,
            runner,
        )
        pairing = pairing_store_module.PairingStore(
            directory / 'telegram_pairing.json'
        )
        pairing.approve(pairing.request(_OWNER, 'Owner'), 'main')
        channel = telegram_channel.TelegramChannel(
            api,
            pairing,
            _NoCommands(),
            notification_formatter.NotificationFormatter([]),
            {
                'main': desk,
            },
            {},
            event_bus,
            'main',
            _store(directory),
        )
        return channel, desk, runner

    def test_paired_photo_reaches_the_desk(self, tmp_path: Path) -> None:
        """A photo with a caption from the owner is downloaded and sent to the desk with it."""
        api = _StandInTelegram(
            {
                'photo-file': _picture(picture_format='JPEG'),
            }
        )
        channel, desk, runner = self._channel(tmp_path, api)
        api.receive(
            telegram_api.IncomingMessage(
                chat_id=_OWNER,
                sender_name='Owner',
                text='What do you make of this chart?',
                files=(
                    telegram_api.IncomingFile(
                        file_id='photo-file',
                        file_name='photo-7.jpg',
                        size_bytes=1000,
                    ),
                ),
            )
        )

        asyncio.run(_run_channel(channel, desk, api, _OWNER))

        attachments = runner.triggers[0].attachments
        assert api.downloads == [
            'photo-file',
        ]
        assert attachments[0].file_name == 'photo-7.jpg'
        assert attachments[0].media_type == 'image/jpeg'
        assert 'What do you make of this chart?' in runner.triggers[0].text

    def test_strangers_files_are_never_downloaded(self, tmp_path: Path) -> None:
        """A file from a chat that is not paired is not fetched; the chat gets a pairing code."""
        api = _StandInTelegram(
            {
                'stranger-file': b'anything',
            }
        )
        channel, desk, runner = self._channel(tmp_path, api)
        api.receive(
            telegram_api.IncomingMessage(
                chat_id=_STRANGER,
                sender_name='Stranger',
                text='',
                files=(
                    telegram_api.IncomingFile(
                        file_id='stranger-file',
                        file_name='payload.pdf',
                        size_bytes=10,
                    ),
                ),
            )
        )

        asyncio.run(_run_channel(channel, desk, api, _STRANGER))

        assert api.downloads == []
        assert runner.triggers == []

    def test_large_file_is_refused_before_download(
        self, tmp_path: Path
    ) -> None:
        """A file Telegram reports as over 20 MB is refused without downloading it."""
        api = _StandInTelegram({})
        channel, desk, runner = self._channel(tmp_path, api)
        api.receive(
            telegram_api.IncomingMessage(
                chat_id=_OWNER,
                sender_name='Owner',
                text='',
                files=(
                    telegram_api.IncomingFile(
                        file_id='big-file',
                        file_name='video.pdf',
                        size_bytes=attachment_store_module.MAXIMUM_BYTES + 1,
                    ),
                ),
            )
        )

        asyncio.run(_run_channel(channel, desk, api, _OWNER))

        assert api.downloads == []
        assert 'larger than 20 MB' in api.sent[0][1]
        assert runner.triggers == []


class _NoCommands:
    """Telegram commands that recognise nothing, so every message goes to the desk."""

    def handles(self, text: str) -> bool:
        """Recognises nothing.

        Args:
            text (str): The message.

        Returns:
            bool: Always False.
        """
        del text
        return False

    def menu(self) -> tuple[tuple[str, str], ...]:
        """Offers no command menu.

        Returns:
            tuple[tuple[str, str], ...]: An empty menu.
        """
        return ()
